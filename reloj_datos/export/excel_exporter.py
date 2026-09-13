"""
export/excel_exporter.py

Sección 9 - Exportación (Excel).
Genera un libro .xlsx con la matriz del Reloj de Datos coloreada con
gradiente continuo (blanco -> rojo, más fuerte cuanto mayor la
frecuencia), los totales, las franjas horarias configuradas (también
coloreadas) y, si corresponde, el ranking de modalidades en una hoja
aparte. Incluye título, subtítulo opcional, rango de fechas analizado,
estadísticas (eventos analizados / excluidos por superar 24 hs) y la
firma del software al pie.
"""
from typing import List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..branding import FIRMA_EXPORTACION
from ..i18n import tr
from ..processing.chart_generator import SeccionRanking
from ..processing.color_scale import color_reloj_rgb
from ..processing.franjas_horarias import ResultadoFranja
from ..processing.reloj_figure import GradientesReloj
from ..processing.reloj_matrix import RelojMatrix, dias_semana

_FILL_ENCABEZADO = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
_FONT_ENCABEZADO = Font(color="FFFFFF", bold=True)


def _fill(color_hex: str) -> PatternFill:
    codigo = color_hex.lstrip("#")
    return PatternFill(start_color=codigo, end_color=codigo, fill_type="solid")


class ExcelExporter:
    def exportar(
        self,
        filepath: str,
        reloj: RelojMatrix,
        titulo: str,
        descripcion_rango: str,
        franjas: List[ResultadoFranja],
        secciones_ranking: Optional[Sequence[SeccionRanking]] = None,
        subtitulo: str = "",
        total_eventos: Optional[int] = None,
        excluidos_24h: int = 0,
        descripcion_filtros: str = "",
        color_reloj: str = "rojo",
        incluir_resumen_ejecutivo: bool = False,
    ) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = tr("excel.hoja_reloj")[:31]

        gradientes = GradientesReloj.desde_datos(reloj, franjas, color_alto=color_reloj_rgb(color_reloj))

        fila = 1
        ws.cell(row=fila, column=1, value=titulo).font = Font(bold=True, size=14)
        fila += 1
        if subtitulo:
            ws.cell(row=fila, column=1, value=subtitulo).font = Font(size=11, color="333333")
            fila += 1
        if descripcion_rango:
            ws.cell(row=fila, column=1, value=descripcion_rango).font = Font(italic=True, size=9, color="555555")
            fila += 1
        if total_eventos is not None:
            texto_stats = tr("reloj.eventos_stats", n=total_eventos, ex=excluidos_24h)
            ws.cell(row=fila, column=1, value=texto_stats).font = Font(size=9, color="666666")
            fila += 1
        if descripcion_filtros:
            ws.cell(row=fila, column=1, value=descripcion_filtros).font = Font(size=8.5, color="777777")
            fila += 1

        fila_encabezado = fila + 1
        for hora in range(24):
            celda = ws.cell(row=fila_encabezado, column=hora + 2, value=f"{hora:02d}:00")
            celda.font = _FONT_ENCABEZADO
            celda.fill = _FILL_ENCABEZADO
        for col, texto in ((26, tr("reloj.frec")), (27, tr("reloj.pct"))):
            celda = ws.cell(row=fila_encabezado, column=col, value=texto)
            celda.font = _FONT_ENCABEZADO
            celda.fill = _FILL_ENCABEZADO

        matriz = reloj.matriz
        total_general = reloj.total_general()
        total_por_dia = reloj.total_por_dia()

        for dia in range(7):
            fila_dia = fila_encabezado + 1 + dia
            ws.cell(row=fila_dia, column=1, value=dias_semana()[dia]).font = Font(bold=True)
            for hora in range(24):
                valor = matriz[dia][hora]
                celda = ws.cell(row=fila_dia, column=hora + 2, value=round(float(valor), 2) if valor else None)
                celda.fill = _fill(gradientes.matriz.color_hex(valor))

            valor_dia = total_por_dia[dia]
            pct_dia = (valor_dia / total_general * 100.0) if total_general else 0.0
            color_dia = _fill(gradientes.totales_dia.color_hex(valor_dia))
            c1 = ws.cell(row=fila_dia, column=26, value=round(float(valor_dia), 2))
            c2 = ws.cell(row=fila_dia, column=27, value=f"{pct_dia:.0f}%")
            c1.fill = color_dia
            c2.fill = color_dia
            c1.font = Font(bold=True)
            c2.font = Font(bold=True)

        fila_totales = fila_encabezado + 8
        ws.cell(row=fila_totales, column=1, value=tr("reloj.frec")).font = Font(bold=True)
        total_por_hora = reloj.total_por_hora()
        for hora in range(24):
            valor = total_por_hora[hora]
            celda = ws.cell(row=fila_totales, column=hora + 2, value=round(float(valor), 2))
            celda.fill = _fill(gradientes.totales_hora.color_hex(valor))
            celda.font = Font(bold=True)
        ws.cell(row=fila_totales, column=26, value=round(float(total_general), 2)).font = Font(bold=True)
        ws.cell(row=fila_totales, column=27, value="100%").font = Font(bold=True)

        fila_franja_frec = fila_totales + 2
        fila_franja_pct = fila_franja_frec + 1
        fila_franja_label = fila_franja_pct + 1
        ws.cell(row=fila_franja_frec, column=1, value=tr("reloj.frec")).font = Font(bold=True)
        ws.cell(row=fila_franja_pct, column=1, value=tr("reloj.pct")).font = Font(bold=True)
        for franja in franjas:
            col_inicio = min(franja.horas) + 2
            color = _fill(gradientes.franjas.color_hex(franja.valor))

            c1 = ws.cell(row=fila_franja_frec, column=col_inicio, value=round(franja.valor, 2))
            c1.fill = color
            c1.font = Font(bold=True)
            c2 = ws.cell(row=fila_franja_pct, column=col_inicio, value=f"{franja.porcentaje:.0f}%")
            c2.fill = color
            c2.font = Font(bold=True)
            c3 = ws.cell(row=fila_franja_label, column=col_inicio, value=franja.etiqueta)
            c3.font = Font(size=8)

        for col in range(1, 28):
            ws.column_dimensions[get_column_letter(col)].width = 8
        ws.column_dimensions["A"].width = 12

        fila_firma = fila_franja_label + 2
        ws.cell(row=fila_firma, column=1, value=FIRMA_EXPORTACION).font = Font(italic=True, size=8, color="999999")

        if secciones_ranking:
            secciones_validas = [(n, r, col) for n, r, col in secciones_ranking if r]
            for numero, (nombre_seccion, ranking, nombre_columna) in enumerate(secciones_validas, start=2):
                nombre_hoja = f"{numero}. {nombre_seccion}"[:31]  # límite de Excel para nombres de hoja
                ws2 = wb.create_sheet(nombre_hoja)
                encabezados = [tr("tabla.posicion"), nombre_columna, tr("tabla.frecuencia"), tr("tabla.porcentaje")]
                for col, texto in enumerate(encabezados, start=1):
                    celda = ws2.cell(row=1, column=col, value=texto)
                    celda.font = _FONT_ENCABEZADO
                    celda.fill = _FILL_ENCABEZADO
                for fila_idx, item in enumerate(ranking, start=2):
                    ws2.cell(row=fila_idx, column=1, value=item.posicion)
                    ws2.cell(row=fila_idx, column=2, value=item.modalidad)
                    ws2.cell(row=fila_idx, column=3, value=item.frecuencia)
                    ws2.cell(row=fila_idx, column=4, value=f"{item.porcentaje:.1f}%")
                ws2.column_dimensions["B"].width = 30
                fila_firma2 = len(ranking) + 3
                ws2.cell(row=fila_firma2, column=1, value=FIRMA_EXPORTACION).font = Font(
                    italic=True, size=8, color="999999"
                )

        if incluir_resumen_ejecutivo:
            secciones_validas = [(n, r, col) for n, r, col in (secciones_ranking or []) if r]
            ws3 = wb.create_sheet(tr("excel.hoja_resumen")[:31])
            ws3.column_dimensions["A"].width = 26
            ws3.column_dimensions["B"].width = 34

            dia_pico = dias_semana()[reloj.dia_maximo()]
            hora_pico = reloj.hora_maxima()
            franja_pico = max(franjas, key=lambda f: f.valor) if franjas else None

            fila_r = 1
            ws3.cell(row=fila_r, column=1, value=tr("graficos.titulo_resumen", titulo=titulo)).font = Font(bold=True, size=14)
            fila_r += 2

            indicadores = [
                (tr("resumen.eventos_analizados"), total_eventos if total_eventos is not None else reloj.eventos_procesados),
                (tr("resumen.dia_pico"), dia_pico),
                (tr("resumen.hora_pico"), tr("resumen.hora_formato", h=hora_pico)),
            ]
            if franja_pico is not None:
                indicadores.append((tr("resumen.franja_pico"), franja_pico.etiqueta))
            for _nombre_seccion, ranking, nombre_columna in secciones_validas:
                indicadores.append((tr("resumen.principal", nombre=nombre_columna.upper()), ranking[0].modalidad))

            for etiqueta, valor in indicadores:
                ws3.cell(row=fila_r, column=1, value=etiqueta).font = Font(bold=True, color="555555")
                ws3.cell(row=fila_r, column=2, value=valor).font = Font(bold=True, size=12)
                fila_r += 1

            for nombre_seccion, ranking, _nombre_columna in secciones_validas:
                fila_r += 1
                ws3.cell(row=fila_r, column=1, value=tr("resumen.top_n", n=3, nombre=nombre_seccion)).font = Font(bold=True)
                fila_r += 1
                for item in ranking[:3]:
                    ws3.cell(row=fila_r, column=1, value=item.modalidad)
                    ws3.cell(row=fila_r, column=2, value=item.frecuencia)
                    fila_r += 1

            fila_r += 1
            ws3.cell(row=fila_r, column=1, value=FIRMA_EXPORTACION).font = Font(italic=True, size=8, color="999999")

        wb.save(filepath)
